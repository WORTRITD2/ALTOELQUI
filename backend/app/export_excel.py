"""Exportación a Excel con el formato de la plantilla original."""
from __future__ import annotations

import datetime as dt
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FUENTE = "Arial"
TITULO = Font(name=FUENTE, size=12, bold=True)
CABECERA = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
NORMAL = Font(name=FUENTE, size=10)
NEGRITA = Font(name=FUENTE, size=10, bold=True)
RELLENO = PatternFill("solid", fgColor="1F3864")
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)
PESOS = '"$"#,##0;("$"#,##0);-'
PORC = "0.00%"


def _hoja_cabecera(ws, titulo: str, obra: dict, fecha: str) -> int:
    ws["A1"] = titulo
    ws["A1"].font = TITULO
    ws["A2"] = f"Obra: {obra.get('nombre', '')}"
    ws["A3"] = f"Licitación: {obra.get('licitacion') or '—'}"
    ws["A4"] = f"Ubicación: {obra.get('ubicacion') or '—'}"
    ws["E2"] = f"Fecha de corte: {fecha}"
    ws["E3"] = f"Contratista: {obra.get('contratista') or '—'}"
    for celda in ("A2", "A3", "A4", "E2", "E3"):
        ws[celda].font = NORMAL
    return 6


def _fila_cabecera(ws, fila: int, columnas: list[tuple[str, int]]) -> None:
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        c = ws.cell(fila, i, titulo)
        c.font = CABECERA
        c.fill = RELLENO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[c.column_letter].width = ancho


def exportar_presupuesto(reporte: dict, obra: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Avance vs presupuesto"

    fila = _hoja_cabecera(ws, "AVANCE REAL SEGÚN PRESUPUESTO", obra, reporte["fecha_corte"])
    columnas = [
        ("Ítem", 12), ("Descripción", 52), ("Unidad", 9), ("Cantidad", 12),
        ("P. Unitario ($)", 14), ("P. Total ($)", 16), ("Cant. avance", 13),
        ("Avance %", 10), ("Monto avance ($)", 16), ("Saldo cantidad", 13),
        ("Saldo ($)", 16), ("Incidencia %", 11),
    ]
    _fila_cabecera(ws, fila, columnas)
    fila += 1

    for f in reporte["filas"]:
        valores = [
            f["codigo"], f["descripcion"], f["unidad"],
            float(f["cantidad"]), float(f["p_unitario"]), float(f["p_total"]),
            float(f["cant_acumulada"]), f["avance_pct"], float(f["monto_avance"]),
            float(f["saldo_cantidad"]), float(f["saldo_monto"]), f["incidencia"],
        ]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(fila, i, v)
            c.font = NORMAL
            c.border = BORDE
            if i in (5, 6, 9, 11):
                c.number_format = PESOS
            if i in (8, 12):
                c.number_format = PORC
        fila += 1

    fila += 1
    contrato, avance = reporte["contrato"], reporte["avance"]
    for etiqueta, clave in (
        ("COSTO DIRECTO NETO", "cd"), ("GASTOS GENERALES", "gg"), ("UTILIDADES", "utilidad"),
        ("SUBTOTAL", "neto"), ("IVA", "iva"), ("TOTAL", "total"),
    ):
        ws.cell(fila, 1, etiqueta).font = NEGRITA
        c1 = ws.cell(fila, 6, float(contrato[clave]))
        c2 = ws.cell(fila, 9, float(avance[clave]))
        for c in (c1, c2):
            c.font = NEGRITA
            c.number_format = PESOS
        fila += 1

    ws.cell(fila, 1, "% AVANCE").font = NEGRITA
    c = ws.cell(fila, 9, reporte["avance_pct"])
    c.font = NEGRITA
    c.number_format = PORC

    # Hoja de capítulos
    ws2 = wb.create_sheet("Capítulos")
    _fila_cabecera(ws2, 1, [("Capítulo", 14), ("Contrato ($)", 18), ("Avance ($)", 18),
                            ("Avance %", 12), ("Incidencia %", 12), ("Partidas", 10)])
    for i, cap in enumerate(reporte["capitulos"], start=2):
        for j, v in enumerate(
            [cap["capitulo"], float(cap["contrato"]), float(cap["avance"]), cap["pct"],
             cap["incidencia"], cap["partidas"]], start=1
        ):
            c = ws2.cell(i, j, v)
            c.font = NORMAL
            if j in (2, 3):
                c.number_format = PESOS
            if j in (4, 5):
                c.number_format = PORC

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exportar_estado_pago(reporte: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"EP N{reporte['numero']}"

    fila = _hoja_cabecera(
        ws, f"ESTADO DE PAGO N°{reporte['numero']}", reporte["obra"], reporte["fecha_corte"]
    )
    ws.cell(fila - 1, 1, f"Período: {reporte['periodo']['desde'] or '—'} a {reporte['periodo']['hasta']}")
    _fila_cabecera(ws, fila, [
        ("Ítem", 12), ("Descripción", 50), ("Unidad", 9), ("Cant. contratada", 14),
        ("Acum. anterior", 14), ("Del período", 13), ("Acum. actual", 14),
        ("% acum.", 10), ("Monto período ($)", 17), ("Monto acum. ($)", 17), ("Saldo ($)", 16),
    ])
    fila += 1
    for d in reporte["detalle"]:
        valores = [
            d["codigo"], d["descripcion"], d["unidad"], float(d["cantidad_contratada"]),
            float(d["cant_anterior"]), float(d["cant_periodo"]), float(d["cant_acumulada"]),
            d["pct_acumulado"], float(d["monto_periodo"]), float(d["monto_acumulado"]),
            float(d["saldo"]),
        ]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(fila, i, v)
            c.font = NORMAL
            c.border = BORDE
            if i in (9, 10, 11):
                c.number_format = PESOS
            if i == 8:
                c.number_format = PORC
        fila += 1

    fila += 1
    pe = reporte["periodo_economico"]
    for etiqueta, clave in (
        ("COSTO DIRECTO", "cd"), ("GASTOS GENERALES", "gg"), ("UTILIDAD", "utilidad"),
        ("NETO", "neto"), ("IVA", "iva"), ("TOTAL DEL PERÍODO", "total"),
        ("Reajuste", "reajuste"), ("Amortización anticipo", "anticipo"),
        ("Retención", "retencion"), ("Multa", "multa"), ("LÍQUIDO A PAGAR", "liquido"),
    ):
        ws.cell(fila, 8, etiqueta).font = NEGRITA
        c = ws.cell(fila, 9, float(pe[clave]))
        c.font = NEGRITA
        c.number_format = PESOS
        fila += 1

    fila += 1
    ws.cell(fila, 1, "Trazabilidad").font = NEGRITA
    fila += 1
    for linea in reporte["pie_trazabilidad"]:
        ws.cell(fila, 1, linea).font = NORMAL
        fila += 1
    for adv in reporte.get("advertencias", []):
        if adv:
            c = ws.cell(fila, 1, f"ADVERTENCIA: {adv}")
            c.font = Font(name=FUENTE, size=10, bold=True, color="C00000")
            fila += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exportar_semanal(reporte: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Avance vs programa"
    ws["A1"] = "AVANCE REAL VS PROGRAMACIÓN — SEMANAS LABORALES HÁBILES"
    ws["A1"].font = TITULO
    ws["A2"] = f"Obra: {reporte['obra']}"
    ws["A3"] = f"Período: {reporte['desde']} a {reporte['hasta']}"

    _fila_cabecera(ws, 5, [
        ("Semana ISO", 12), ("Inicio", 12), ("Fin", 12), ("Días hábiles", 12),
        ("Feriados", 26), ("Programado período", 17), ("Real período", 14),
        ("Programado acum.", 17), ("Real acum.", 13), ("Δ (p.p.)", 11), ("SPI", 9),
        ("Monto ejecutado ($)", 18), ("Semáforo", 12),
    ])
    fila = 6
    for s in reporte["semanas"]:
        valores = [
            s["iso"], s["inicio"], s["fin"], s["dias_habiles"], ", ".join(s["feriados"]),
            s["programado_periodo"], s["real_periodo"], s["programado_acumulado"],
            s["real_acumulado"], s["desviacion_pp"], s["spi"],
            float(s["monto_ejecutado_periodo"]), s["semaforo"],
        ]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(fila, i, v)
            c.font = NORMAL
            c.border = BORDE
            if i in (6, 7, 8, 9):
                c.number_format = PORC
            if i == 12:
                c.number_format = PESOS
        fila += 1

    ws2 = wb.create_sheet("Partidas críticas")
    _fila_cabecera(ws2, 1, [
        ("Ítem", 12), ("Descripción", 50), ("Unidad", 9), ("Cant. programada", 15),
        ("Cant. ejecutada", 15), ("Avance programado", 16), ("Avance real", 13),
        ("Atraso (p.p.)", 13), ("Monto de atraso ($)", 18), ("Semáforo", 12),
    ])
    for i, p in enumerate(reporte["partidas_criticas"], start=2):
        valores = [
            p["codigo"], p["descripcion"], p["unidad"], float(p["cantidad_programada"]),
            float(p["cantidad_ejecutada"]), p["avance_programado"], p["avance_real"],
            p["atraso_pp"], float(p["monto_atraso"]), p["semaforo"],
        ]
        for j, v in enumerate(valores, start=1):
            c = ws2.cell(i, j, v)
            c.font = NORMAL
            if j in (6, 7):
                c.number_format = PORC
            if j == 9:
                c.number_format = PESOS

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exportar_anexo4(anexo: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Anexo 4"
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 50
    for col in "CDEF":
        ws.column_dimensions[col].width = 14

    ws["A1"] = "Anexo N°4"
    ws["A1"].font = TITULO
    ws["A3"] = "ANÁLISIS DE PRECIOS UNITARIOS"
    ws["A3"].font = NEGRITA
    ws["A5"] = f"NOMBRE DE LA OBRA: {anexo['obra']}"
    ws["A6"] = f"ID: {anexo.get('licitacion') or '—'}"
    ws["A7"] = f"CONTRATISTA: {anexo.get('contratista') or '—'}"

    p = anexo["partida"]
    ws["A9"], ws["B9"] = "PARTIDA", p["descripcion"]
    ws["C9"], ws["D9"], ws["F9"] = "ITEM :", "UNIDAD", "PRECIO"
    ws["C10"], ws["D10"], ws["F10"] = p["codigo"], p["unidad"], "UNITARIO"
    ws["F11"] = float(anexo["precio_unitario"])
    ws["F11"].number_format = PESOS
    ws["F11"].font = NEGRITA

    fila = 13
    for titulo, clave, letra in (
        ("MATERIALES (A)", "materiales", "A"),
        ("MANO DE OBRA (B)", "mano_obra", "B"),
        ("EQUIPOS Y HERRAMIENTAS (C)", "equipos", "C"),
    ):
        ws.cell(fila, 1, titulo).font = NEGRITA
        fila += 1
        _fila_cabecera(ws, fila, [("Nº", 6), ("ITEM", 50), ("UNIDAD", 14),
                                  ("CANTIDAD", 14), ("PRECIO UNITARIO", 14), ("VALOR TOTAL", 14)])
        fila += 1
        for i, item in enumerate(anexo[clave]["items"], start=1):
            for j, v in enumerate(
                [i, item["descripcion"], item["unidad"], float(item["cantidad"] or 0),
                 float(item["precio_unitario"] or 0), float(item["valor_total"] or 0)], start=1
            ):
                c = ws.cell(fila, j, v)
                c.font = NORMAL
                if j in (5, 6):
                    c.number_format = PESOS
            fila += 1
        if clave == "mano_obra":
            ws.cell(fila, 2, "LEYES SOCIALES").font = NEGRITA
            ws.cell(fila, 3, f"{anexo['mano_obra']['pct_leyes_sociales'] * 100:.0f}%")
            c = ws.cell(fila, 6, float(anexo["mano_obra"]["leyes_sociales"]))
            c.number_format = PESOS
            fila += 1
        ws.cell(fila, 5, "TOTAL").font = NEGRITA
        c = ws.cell(fila, 6, float(anexo[clave]["total"]))
        c.font = NEGRITA
        c.number_format = PESOS
        fila += 2

    ws.cell(fila + 2, 2, "____________________________________")
    ws.cell(fila + 3, 2, "REPRESENTANTE LEGAL")
    ws.cell(fila + 4, 2, anexo.get("contratista") or "")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
