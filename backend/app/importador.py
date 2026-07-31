"""Importador de la plantilla de obra (§7 del prompt).

Lee el .xlsx tal como lo entrega la oficina técnica y **reporta cada anomalía**
en vez de silenciarla. Las ocho validaciones corresponden a defectos reales
detectados en la plantilla de referencia.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    APU,
    APURecurso,
    Obra,
    Partida,
    Presupuesto,
    ValidacionImportacion,
)

TOLERANCIA_PESOS = Decimal("1")

UNIDADES = {
    "m²": "m2", "m2": "m2", "M2": "m2", "mt2": "m2",
    "m": "m", "ml": "m", "m. lineal": "m", "m lineal": "m", "mt": "m",
    "un": "un", "unidad": "un", "u": "un", "c/u": "un",
    "gl": "gl", "global": "gl",
    "m3": "m3", "m³": "m3",
    "kg": "kg", "día": "dia", "dia": "dia", "hr": "hr", "h": "hr",
}

ETIQUETAS_CIERRE = {
    "COSTO DIRECTO NETO": "cd",
    "GASTOS GENERALES": "gg",
    "UTILIDADES": "utilidad",
    "SUBTOTAL": "neto",
    "IVA": "iva",
    "TOTAL": "total",
}


@dataclass
class Hallazgo:
    regla: str
    severidad: str  # INFO | ADVERTENCIA | ERROR
    mensaje: str
    ubicacion: str | None = None
    valor_original: str | None = None
    valor_corregido: str | None = None


@dataclass
class ResultadoImportacion:
    presupuesto: Presupuesto | None
    hallazgos: list[Hallazgo] = field(default_factory=list)
    partidas_medibles: int = 0
    agrupadores: int = 0
    incluidas_en_gg: int = 0
    apus: int = 0
    totales_archivo: dict[str, Decimal] = field(default_factory=dict)
    totales_calculados: dict[str, Decimal] = field(default_factory=dict)
    reutilizado: bool = False

    @property
    def errores(self) -> list[Hallazgo]:
        return [h for h in self.hallazgos if h.severidad == "ERROR"]

    def resumen(self) -> dict:
        return {
            "presupuesto_id": self.presupuesto.id if self.presupuesto else None,
            "reutilizado": self.reutilizado,
            "partidas_medibles": self.partidas_medibles,
            "agrupadores": self.agrupadores,
            "incluidas_en_gg": self.incluidas_en_gg,
            "no_medibles": self.agrupadores + self.incluidas_en_gg,
            "apus": self.apus,
            "hallazgos": len(self.hallazgos),
            "errores": len(self.errores),
            "por_regla": {
                r: len([h for h in self.hallazgos if h.regla == r])
                for r in sorted({h.regla for h in self.hallazgos})
            },
            "totales_archivo": {k: str(v) for k, v in self.totales_archivo.items()},
            "totales_calculados": {k: str(v) for k, v in self.totales_calculados.items()},
        }


# --------------------------------------------------------------------------
# Utilidades de lectura
# --------------------------------------------------------------------------
def _dec(valor) -> Decimal | None:
    if valor is None or isinstance(valor, (dt.datetime, dt.date)):
        return None
    if isinstance(valor, str):
        limpio = valor.strip().replace("$", "").replace(".", "").replace(",", ".")
        if not limpio or not re.match(r"^-?\d+(\.\d+)?$", limpio):
            return None
        valor = limpio
    try:
        return Decimal(str(valor))
    except (InvalidOperation, ValueError):
        return None


def normalizar_unidad(texto) -> tuple[str | None, str | None]:
    if texto is None:
        return None, None
    original = str(texto).strip()
    return UNIDADES.get(original.lower(), UNIDADES.get(original, original.lower())), original


def reconstruir_codigo(valor: dt.datetime | dt.date) -> str:
    """Excel leyó `2.6.10` como fecha m.d.yy → (2010, 2, 6). Se revierte."""
    return f"{valor.month}.{valor.day}.{valor.year % 100}"


def _codigo_de(celda) -> tuple[str | None, str | None]:
    """Devuelve (codigo, codigo_original_si_estaba_corrupto)."""
    if celda is None:
        return None, None
    if isinstance(celda, (dt.datetime, dt.date)):
        return reconstruir_codigo(celda), celda.strftime("%Y-%m-%d")
    if isinstance(celda, float) and celda.is_integer():
        return str(int(celda)), None
    return str(celda).strip(), None


def _nivel(codigo: str) -> int:
    return len([p for p in str(codigo).split(".") if p != ""])


def _padre(codigo: str) -> str | None:
    partes = str(codigo).split(".")
    return ".".join(partes[:-1]) if len(partes) > 1 else None


# --------------------------------------------------------------------------
# Importación del itemizado
# --------------------------------------------------------------------------
def _hoja_presupuesto(wb, nombre: str | None):
    if nombre and nombre in wb.sheetnames:
        return wb[nombre]
    for hoja in wb.sheetnames:
        if hoja.strip().lower().startswith("avance"):
            return wb[hoja]
    return wb[wb.sheetnames[0]]


def _fila_encabezado(ws) -> int:
    for r in range(1, 30):
        if str(ws.cell(r, 1).value or "").strip().lower() in ("ítem", "item"):
            return r
    return 4


def importar_itemizado(ws, hallazgos: list[Hallazgo]) -> tuple[list[dict], dict[str, Decimal]]:
    encabezado = _fila_encabezado(ws)
    filas: list[dict] = []
    totales: dict[str, Decimal] = {}
    capitulo_actual: str | None = None
    orden = 0

    for r in range(encabezado + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        etiqueta = str(a or "").strip().upper()

        if etiqueta in ETIQUETAS_CIERRE:
            valor = _dec(ws.cell(r, 6).value)
            if valor is not None:
                totales[ETIQUETAS_CIERRE[etiqueta]] = valor
            continue

        # Fila SUBTOTAL de capítulo, rotulada en C o D (regla 4)
        rotulo_sub = next(
            (
                str(ws.cell(r, col).value).strip().upper()
                for col in (3, 4, 5)
                if str(ws.cell(r, col).value or "").strip().upper() == "SUBTOTAL"
            ),
            None,
        )
        if rotulo_sub and a is None:
            if _dec(ws.cell(r, 6).value) is None:
                hallazgos.append(
                    Hallazgo(
                        "SUBTOTAL_SIN_FORMULA",
                        "ADVERTENCIA",
                        "Fila rotulada SUBTOTAL sin valor: se recalcula desde las partidas.",
                        ubicacion=f"{ws.title}!{r}",
                    )
                )
            continue

        if a is None and b is None:
            continue

        codigo, corrupto = _codigo_de(a)
        if codigo is None:
            continue
        if corrupto:
            hallazgos.append(
                Hallazgo(
                    "CODIGO_FECHA",
                    "ADVERTENCIA",
                    f"Código de ítem convertido a fecha por Excel; reconstruido como {codigo}.",
                    ubicacion=f"{ws.title}!A{r}",
                    valor_original=corrupto,
                    valor_corregido=codigo,
                )
            )

        descripcion = str(b or "").strip()
        unidad, unidad_original = normalizar_unidad(c)
        cantidad = _dec(ws.cell(r, 4).value)
        p_unitario = _dec(ws.cell(r, 5).value)
        p_total = _dec(ws.cell(r, 6).value)
        cant_avance = _dec(ws.cell(r, 7).value)

        incluida_gg = bool(c and "INCLUIDO" in str(c).upper())
        if incluida_gg:
            unidad, unidad_original = None, str(c).strip()
            hallazgos.append(
                Hallazgo(
                    "INCLUIDA_EN_GG",
                    "INFO",
                    f"{codigo} {descripcion}: marcada «incluido en gastos generales», "
                    "se importa como informativa y no pondera en el avance.",
                    ubicacion=f"{ws.title}!{r}",
                )
            )

        es_agrupador = not incluida_gg and (cantidad is None or p_unitario is None)

        if unidad_original and unidad and unidad_original.lower() != unidad:
            hallazgos.append(
                Hallazgo(
                    "UNIDAD_NO_NORMALIZADA",
                    "INFO",
                    f"Unidad «{unidad_original}» normalizada a «{unidad}».",
                    ubicacion=f"{ws.title}!C{r}",
                    valor_original=unidad_original,
                    valor_corregido=unidad,
                )
            )

        if not es_agrupador and not incluida_gg and cantidad and p_unitario:
            esperado = cantidad * p_unitario
            if p_total is None or abs(p_total - esperado) > TOLERANCIA_PESOS:
                hallazgos.append(
                    Hallazgo(
                        "PTOTAL_DESCUADRADO",
                        "ADVERTENCIA",
                        f"{codigo}: P. Total del archivo {p_total} ≠ cantidad × P. Unitario "
                        f"({esperado}). Se recalcula.",
                        ubicacion=f"{ws.title}!F{r}",
                        valor_original=str(p_total),
                        valor_corregido=str(esperado),
                    )
                )
                p_total = esperado

        if _nivel(codigo) == 1:
            capitulo_actual = codigo

        orden += 1
        filas.append(
            {
                "codigo": codigo,
                "codigo_padre": _padre(codigo),
                "nivel": _nivel(codigo),
                "descripcion": descripcion,
                "unidad": unidad,
                "unidad_original": unidad_original,
                "cantidad": cantidad,
                "p_unitario": p_unitario,
                "p_total": p_total,
                "es_agrupador": es_agrupador,
                "es_incluida_en_gg": incluida_gg,
                "orden": orden,
                "fila_origen": r,
                "cant_avance": cant_avance,
                "capitulo": capitulo_actual,
            }
        )

    # Regla 6: celdas sueltas fuera de la tabla
    for r in range(encabezado + 1, ws.max_row + 1):
        for col in range(11, min(ws.max_column, 26) + 1):
            valor = ws.cell(r, col).value
            if valor not in (None, ""):
                hallazgos.append(
                    Hallazgo(
                        "CELDA_HUERFANA",
                        "ADVERTENCIA",
                        f"Valor fuera de la tabla, ignorado: {valor!r}.",
                        ubicacion=f"{ws.title}!{ws.cell(r, col).coordinate}",
                        valor_original=str(valor),
                    )
                )
    return filas, totales


# --------------------------------------------------------------------------
# Importación de los APU (hoja P.U)
# --------------------------------------------------------------------------
SECCIONES = {"MATERIALES": "MATERIAL", "MANO DE OBRA": "MANO_OBRA", "MAQUINARIA": "EQUIPO",
             "EQUIPOS": "EQUIPO"}


RE_CODIGO = re.compile(r"^\d+(\.\d+)+$")


def _titulo_bloque(ws, inicio: int, fin: int) -> tuple[str | None, str]:
    """El código de la partida aparece a veces en B («2.1.1.1 Retiro…») y a veces
    en A con la descripción en B. Se aceptan ambos."""
    for r in range(inicio, min(inicio + 6, fin + 1)):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and RE_CODIGO.match(a.strip()):
            return a.strip(), str(b or "").strip()
        if isinstance(b, str) and re.match(r"^\d+(\.\d+)+", b.strip()):
            texto = b.strip()
            codigo = texto.split()[0]
            return codigo, texto[len(codigo):].strip()
    return None, ""


def importar_apus(ws, hallazgos: list[Hallazgo], orden_codigos: list[str] | None = None) -> dict[str, dict]:
    """Devuelve {codigo_partida: {rendimiento, unidad_pago, recursos, resumen, fila}}."""
    inicios = [
        r
        for r in range(1, ws.max_row + 1)
        if str(ws.cell(r, 1).value or "").strip().lower().startswith("rendimiento")
    ]
    apus: dict[str, dict] = {}
    sin_codigo: list[dict] = []
    secuencia: list[str | None] = []

    for idx, inicio in enumerate(inicios):
        fin = inicios[idx + 1] - 1 if idx + 1 < len(inicios) else ws.max_row
        rendimiento = _dec(ws.cell(inicio, 2).value)
        unidad_pago = ws.cell(inicio + 1, 2).value

        codigo, descripcion = _titulo_bloque(ws, inicio, fin)

        recursos: list[dict] = []
        resumen: dict[str, Decimal] = {}
        seccion: str | None = None
        orden = 0

        for r in range(inicio, fin + 1):
            a = str(ws.cell(r, 1).value or "").strip().upper()
            e = str(ws.cell(r, 5).value or "").strip().upper()

            for nombre, tipo in SECCIONES.items():
                if a.startswith(nombre):
                    seccion = tipo
                    break
            if a.startswith("RESUMEN"):
                seccion = None
            if e in ("COSTO DIRECTO", "GASTOS GENERALES", "UTILIDAD", "SUB-TOTAL NETO",
                     "PRECIO UNITARIO") or e.startswith("IVA"):
                valor = _dec(ws.cell(r, 6).value)
                if valor is not None:
                    clave = {"COSTO DIRECTO": "cd", "GASTOS GENERALES": "gg", "UTILIDAD": "utilidad",
                             "SUB-TOTAL NETO": "neto", "PRECIO UNITARIO": "precio_unitario"}.get(e, "iva")
                    resumen[clave] = valor
                continue

            if seccion is None:
                continue
            numero = ws.cell(r, 1).value
            descripcion_rec = ws.cell(r, 2).value
            es_item = isinstance(numero, (int, float)) or (
                isinstance(numero, str) and numero.strip().isdigit()
            )
            if not es_item or not descripcion_rec:
                continue
            orden += 1
            unidad, _ = normalizar_unidad(ws.cell(r, 3).value)
            recursos.append(
                {
                    "tipo": seccion,
                    "descripcion": str(descripcion_rec).strip(),
                    "unidad": unidad,
                    "cantidad": _dec(ws.cell(r, 4).value),
                    "precio_unitario": _dec(ws.cell(r, 5).value),
                    "subtotal": _dec(ws.cell(r, 6).value),
                    "orden": orden,
                }
            )

        # Regla 5: IVA del resumen inconsistente
        if "neto" in resumen and "iva" in resumen:
            esperado = resumen["neto"] * Decimal("0.19")
            if abs(resumen["iva"] - esperado) > Decimal("2"):
                hallazgos.append(
                    Hallazgo(
                        "APU_IVA_INCONSISTENTE",
                        "ERROR",
                        f"APU {codigo}: IVA del resumen {resumen['iva']} ≠ neto × 19% "
                        f"({esperado.quantize(Decimal('1'))}). Se recalcula desde los parámetros.",
                        ubicacion=f"{ws.title}!{inicio}",
                        valor_original=str(resumen["iva"]),
                        valor_corregido=str(esperado.quantize(Decimal("1"))),
                    )
                )

        bloque = {
            "codigo": codigo,
            "descripcion": descripcion,
            "rendimiento": rendimiento,
            "unidad_pago": (str(unidad_pago).strip() if unidad_pago else None),
            "recursos": recursos,
            "resumen": resumen,
            "fila": inicio,
        }
        secuencia.append(codigo)
        if codigo:
            apus[codigo] = bloque
        else:
            sin_codigo.append({"bloque": bloque, "posicion": len(secuencia) - 1})

    # Inferencia por posición para bloques sin código legible (§7 punto 1, mismo criterio)
    for pendiente in sin_codigo:
        bloque, pos = pendiente["bloque"], pendiente["posicion"]
        previo = next((c for c in reversed(secuencia[:pos]) if c), None)
        siguiente = next((c for c in secuencia[pos + 1:] if c), None)
        candidato = None
        if orden_codigos and previo in orden_codigos and siguiente in orden_codigos:
            i, j = orden_codigos.index(previo), orden_codigos.index(siguiente)
            entre = [c for c in orden_codigos[i + 1: j] if c not in apus]
            if len(entre) == 1:
                candidato = entre[0]
        if candidato:
            bloque["codigo"] = candidato
            apus[candidato] = bloque
            hallazgos.append(
                Hallazgo(
                    "APU_CODIGO_INFERIDO",
                    "ADVERTENCIA",
                    f"Bloque de APU sin código legible; inferido como {candidato} por su posición "
                    f"entre {previo} y {siguiente}. Requiere confirmación.",
                    ubicacion=f"{ws.title}!{bloque['fila']}",
                    valor_corregido=candidato,
                )
            )
        else:
            hallazgos.append(
                Hallazgo(
                    "APU_SIN_CODIGO",
                    "ADVERTENCIA",
                    "Bloque de APU sin código de partida identificable ni inferible por posición.",
                    ubicacion=f"{ws.title}!{bloque['fila']}",
                )
            )
    return apus


# --------------------------------------------------------------------------
# Mediciones (columna «Cant avance» de una hoja de corte)
# --------------------------------------------------------------------------
def leer_mediciones(ruta: str, hoja: str) -> tuple[dt.date | None, dict[str, Decimal]]:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb[hoja]
    fecha = None
    for r in range(1, 6):
        for col in range(1, 10):
            valor = ws.cell(r, col).value
            if isinstance(valor, str) and valor.strip().lower().startswith("fecha"):
                texto = valor.split(":", 1)[-1].strip()
                for formato in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
                    try:
                        fecha = dt.datetime.strptime(texto, formato).date()
                        break
                    except ValueError:
                        continue
            elif isinstance(valor, (dt.datetime, dt.date)) and fecha is None and r <= 3:
                fecha = valor.date() if isinstance(valor, dt.datetime) else valor

    encabezado = _fila_encabezado(ws)
    mediciones: dict[str, Decimal] = {}
    for r in range(encabezado + 1, ws.max_row + 1):
        codigo, _ = _codigo_de(ws.cell(r, 1).value)
        if not codigo:
            continue
        cantidad = _dec(ws.cell(r, 7).value)
        if cantidad is not None:
            mediciones[codigo] = cantidad
    return fecha, mediciones


# --------------------------------------------------------------------------
# Orquestación
# --------------------------------------------------------------------------
def hash_archivo(ruta: str) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as fh:
        for bloque in iter(lambda: fh.read(65536), b""):
            h.update(bloque)
    return h.hexdigest()[:32]


def importar_plantilla(
    db: Session,
    obra: Obra,
    ruta: str,
    hoja_presupuesto: str | None = None,
    hoja_pu: str = "P.U",
    pct_leyes_sociales: Decimal = Decimal("0.5"),
) -> ResultadoImportacion:
    firma = hash_archivo(ruta)
    origen = f"{ruta}#{firma}"

    ya = db.scalars(
        select(Presupuesto).where(
            Presupuesto.obra_id == obra.id, Presupuesto.origen_archivo == origen
        )
    ).first()
    if ya:
        return ResultadoImportacion(presupuesto=ya, reutilizado=True)

    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = _hoja_presupuesto(wb, hoja_presupuesto)
    hallazgos: list[Hallazgo] = []

    filas, totales_archivo = importar_itemizado(ws, hallazgos)
    orden_codigos = [f["codigo"] for f in filas if not f["es_agrupador"] and not f["es_incluida_en_gg"]]
    apus = (
        importar_apus(wb[hoja_pu], hallazgos, orden_codigos) if hoja_pu in wb.sheetnames else {}
    )

    ultima = db.scalars(
        select(Presupuesto)
        .where(Presupuesto.obra_id == obra.id)
        .order_by(Presupuesto.version.desc())
    ).first()
    presupuesto = Presupuesto(
        obra_id=obra.id,
        version=(ultima.version + 1) if ultima else 1,
        estado="BORRADOR",
        vigente_desde=obra.fecha_inicio,
        origen_archivo=origen,
    )
    db.add(presupuesto)
    db.flush()

    creadas: dict[str, Partida] = {}
    cd_calculado = Decimal("0")
    medibles = agrupadores = incluidas = 0

    for f in filas:
        partida = Partida(
            presupuesto_id=presupuesto.id,
            codigo=f["codigo"],
            codigo_padre=f["codigo_padre"],
            nivel=f["nivel"],
            descripcion=f["descripcion"],
            unidad=f["unidad"],
            unidad_original=f["unidad_original"],
            cantidad=f["cantidad"],
            p_unitario=f["p_unitario"],
            p_total=f["p_total"],
            es_agrupador=f["es_agrupador"],
            es_incluida_en_gg=f["es_incluida_en_gg"],
            orden=f["orden"],
            fila_origen=f["fila_origen"],
        )
        db.add(partida)
        db.flush()
        creadas[f["codigo"]] = partida
        if partida.es_medible:
            medibles += 1
            cd_calculado += partida.p_total or Decimal("0")
        elif f["es_agrupador"]:
            agrupadores += 1
        if f["es_incluida_en_gg"]:
            incluidas += 1

    # APU + regla 7 (cuadratura APU ↔ itemizado)
    creados_apu = 0
    for codigo, datos in apus.items():
        partida = creadas.get(codigo)
        if partida is None:
            hallazgos.append(
                Hallazgo(
                    "APU_SIN_PARTIDA",
                    "ADVERTENCIA",
                    f"APU {codigo} sin partida correspondiente en el itemizado.",
                    ubicacion=f"{hoja_pu}!{datos['fila']}",
                )
            )
            continue
        apu = APU(
            partida_id=partida.id,
            rendimiento=datos["rendimiento"],
            unidad_pago=datos["unidad_pago"],
            hoja_origen=hoja_pu,
            fila_origen=datos["fila"],
        )
        db.add(apu)
        db.flush()
        for r in datos["recursos"]:
            db.add(APURecurso(apu_id=apu.id, **r))
        db.flush()
        creados_apu += 1

        cd_apu = sum(
            (
                (r["subtotal"] or Decimal("0"))
                * ((Decimal("1") + pct_leyes_sociales) if r["tipo"] == "MANO_OBRA" else Decimal("1"))
                for r in datos["recursos"]
            ),
            Decimal("0"),
        )
        if partida.p_unitario is not None and abs(cd_apu - partida.p_unitario) > TOLERANCIA_PESOS:
            hallazgos.append(
                Hallazgo(
                    "APU_NO_CUADRA",
                    "ERROR",
                    f"Partida {codigo}: costo directo del APU {cd_apu.quantize(Decimal('1'))} ≠ "
                    f"P. Unitario del itemizado {partida.p_unitario}. Diferencia "
                    f"{(cd_apu - partida.p_unitario).quantize(Decimal('1'))}.",
                    ubicacion=f"{hoja_pu}!{datos['fila']}",
                    valor_original=str(partida.p_unitario),
                    valor_corregido=str(cd_apu.quantize(Decimal("1"))),
                )
            )

    for codigo, partida in creadas.items():
        if partida.es_medible and codigo not in apus:
            hallazgos.append(
                Hallazgo(
                    "PARTIDA_SIN_APU",
                    "INFO",
                    f"Partida medible {codigo} sin APU en la hoja {hoja_pu}.",
                    ubicacion=f"{ws.title}!{partida.fila_origen}",
                )
            )

    # Regla 8: cuadratura de totales
    totales_calculados = {"cd": cd_calculado}
    if "cd" in totales_archivo:
        diferencia = abs(cd_calculado - totales_archivo["cd"])
        if diferencia > TOLERANCIA_PESOS:
            hallazgos.append(
                Hallazgo(
                    "TOTALES_NO_CUADRAN",
                    "ERROR",
                    f"Costo directo calculado {cd_calculado} ≠ del archivo {totales_archivo['cd']} "
                    f"(diferencia {diferencia}).",
                    ubicacion=f"{ws.title}!F",
                    valor_original=str(totales_archivo["cd"]),
                    valor_corregido=str(cd_calculado),
                )
            )
        else:
            hallazgos.append(
                Hallazgo(
                    "TOTALES_CUADRAN",
                    "INFO",
                    f"Costo directo calculado coincide con el archivo: {cd_calculado} "
                    f"(diferencia {diferencia}).",
                )
            )

    for h in hallazgos:
        db.add(
            ValidacionImportacion(
                presupuesto_id=presupuesto.id,
                regla=h.regla,
                severidad=h.severidad,
                ubicacion=h.ubicacion,
                mensaje=h.mensaje,
                valor_original=h.valor_original,
                valor_corregido=h.valor_corregido,
            )
        )

    for clave, columna in (
        ("cd", "total_cd"), ("gg", "total_gg"), ("utilidad", "total_utilidad"),
        ("neto", "total_neto"), ("iva", "total_iva"), ("total", "total_contrato"),
    ):
        if clave in totales_archivo:
            setattr(presupuesto, columna, totales_archivo[clave])

    presupuesto.estado = "VALIDADO" if not [h for h in hallazgos if h.severidad == "ERROR"] else "BORRADOR"
    db.commit()

    return ResultadoImportacion(
        presupuesto=presupuesto,
        hallazgos=hallazgos,
        partidas_medibles=medibles,
        agrupadores=agrupadores,
        incluidas_en_gg=incluidas,
        apus=creados_apu,
        totales_archivo=totales_archivo,
        totales_calculados=totales_calculados,
    )
